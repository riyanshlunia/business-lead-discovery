"""
FastAPI router for Social Lead Finder.
Provides endpoints for searches, leads, analytics, saved searches, and exports.
"""
from __future__ import annotations

import io
import csv
import json

from fastapi import APIRouter, BackgroundTasks, Depends, HTTPException, status
from fastapi.responses import Response, StreamingResponse
from sqlalchemy import func, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.database.session import get_session
from app.models.entities import JobStatus, SavedSearch, SocialLead, SocialSearch
from app.schemas.social import (
    AnalyticsRead,
    SavedSearchCreate,
    SavedSearchRead,
    SocialLeadRead,
    SocialSearchCreate,
    SocialSearchRead,
)
from app.services.social_lead_pipeline import SocialLeadPipeline

router = APIRouter(prefix="/social-searches", tags=["social"])

_pipeline = SocialLeadPipeline()


# ── Social Searches ──────────────────────────────────────────────────────────

@router.post("", response_model=SocialSearchRead, status_code=status.HTTP_202_ACCEPTED)
async def create_social_search(
    payload: SocialSearchCreate,
    background_tasks: BackgroundTasks,
    session: AsyncSession = Depends(get_session),
) -> SocialSearchRead:
    search = SocialSearch(
        industry=payload.industry,
        location=payload.location,
        keywords=payload.keywords,
        sources=payload.sources,
        filters=payload.filters,
        target_limit=max(1, payload.limit) if payload.limit else 20,
        status=JobStatus.queued,
    )
    session.add(search)
    await session.commit()
    await session.refresh(search)

    background_tasks.add_task(_pipeline.run, search.id)
    return SocialSearchRead.model_validate(search, from_attributes=True)


@router.get("", response_model=list[SocialSearchRead])
async def list_social_searches(
    session: AsyncSession = Depends(get_session),
) -> list[SocialSearchRead]:
    result = await session.execute(
        select(SocialSearch).order_by(SocialSearch.created_at.desc()).limit(50)
    )
    searches = result.scalars().all()
    return [SocialSearchRead.model_validate(s, from_attributes=True) for s in searches]


@router.get("/analytics", response_model=AnalyticsRead)
async def get_analytics(
    session: AsyncSession = Depends(get_session),
) -> AnalyticsRead:
    """Aggregate analytics across all social searches and leads."""

    # Counts
    search_count_result = await session.execute(select(func.count()).select_from(SocialSearch))
    total_searches = search_count_result.scalar() or 0

    lead_count_result = await session.execute(select(func.count()).select_from(SocialLead))
    total_leads = lead_count_result.scalar() or 0

    qualified_result = await session.execute(
        select(func.count()).select_from(SocialLead).where(SocialLead.lead_score >= 70)
    )
    qualified_leads = qualified_result.scalar() or 0

    avg_score_result = await session.execute(select(func.avg(SocialLead.lead_score)))
    avg_lead_score = round(float(avg_score_result.scalar() or 0), 1)

    avg_conf_result = await session.execute(select(func.avg(SocialLead.confidence_score)))
    avg_confidence = round(float(avg_conf_result.scalar() or 0), 1)

    # Platform distribution
    platform_result = await session.execute(
        select(SocialLead.source_platform, func.count().label("cnt"))
        .group_by(SocialLead.source_platform)
        .order_by(func.count().desc())
        .limit(15)
    )
    platform_distribution = {row[0]: row[1] for row in platform_result}

    # Status distribution of searches
    status_result = await session.execute(
        select(SocialSearch.status, func.count().label("cnt"))
        .group_by(SocialSearch.status)
    )
    status_distribution = {str(row[0].value if hasattr(row[0], "value") else row[0]): row[1] for row in status_result}

    # Top industries from searches
    industry_result = await session.execute(
        select(SocialSearch.industry, func.count().label("cnt"))
        .group_by(SocialSearch.industry)
        .order_by(func.count().desc())
        .limit(10)
    )
    top_industries = [{"industry": row[0], "count": row[1]} for row in industry_result]

    return AnalyticsRead(
        total_searches=total_searches,
        total_leads=total_leads,
        qualified_leads=qualified_leads,
        avg_lead_score=avg_lead_score,
        avg_confidence=avg_confidence,
        platform_distribution=platform_distribution,
        status_distribution=status_distribution,
        top_industries=top_industries,
    )


@router.get("/{search_id}", response_model=SocialSearchRead)
async def get_social_search(
    search_id: int,
    session: AsyncSession = Depends(get_session),
) -> SocialSearchRead:
    search = await session.get(SocialSearch, search_id)
    if not search:
        raise HTTPException(status_code=404, detail="Search not found")
    return SocialSearchRead.model_validate(search, from_attributes=True)


@router.get("/{search_id}/leads", response_model=list[SocialLeadRead])
async def list_social_leads(
    search_id: int,
    sort: str = "lead_score",
    platform: str | None = None,
    min_score: int | None = None,
    has_email: bool | None = None,
    has_phone: bool | None = None,
    has_website: bool | None = None,
    session: AsyncSession = Depends(get_session),
) -> list[SocialLeadRead]:
    stmt = select(SocialLead).where(SocialLead.search_id == search_id)

    # Filtering
    if platform:
        stmt = stmt.where(SocialLead.source_platform == platform)
    if min_score is not None:
        stmt = stmt.where(SocialLead.lead_score >= min_score)
    if has_email is not None:
        stmt = stmt.where(SocialLead.has_email == has_email)
    if has_phone is not None:
        stmt = stmt.where(SocialLead.has_phone == has_phone)
    if has_website is not None:
        stmt = stmt.where(SocialLead.has_website == has_website)

    # Sorting
    sort_col = {
        "lead_score": SocialLead.lead_score.desc(),
        "confidence_score": SocialLead.confidence_score.desc(),
        "digital_score": SocialLead.digital_score.desc(),
        "name": SocialLead.name.asc(),
        "created_at": SocialLead.created_at.desc(),
    }.get(sort, SocialLead.lead_score.desc())

    stmt = stmt.order_by(sort_col)
    result = await session.execute(stmt)
    leads = result.scalars().all()
    return [SocialLeadRead.model_validate(lead, from_attributes=True) for lead in leads]


# ── Exports ──────────────────────────────────────────────────────────────────

def _build_lead_row(lead: SocialLead) -> dict:
    return {
        "name": lead.name,
        "source_platform": lead.source_platform,
        "website": lead.website or "",
        "email": lead.email or "",
        "phone": lead.phone or "",
        "linkedin": lead.linkedin or "",
        "facebook": lead.facebook or "",
        "instagram": lead.instagram or "",
        "twitter": lead.twitter or "",
        "github": lead.github or "",
        "youtube": lead.youtube or "",
        "industry": lead.industry or "",
        "city": lead.city or "",
        "state": lead.state or "",
        "country": lead.country or "",
        "employee_estimate": lead.employee_estimate or "",
        "lead_score": lead.lead_score,
        "digital_score": lead.digital_score,
        "confidence_score": lead.confidence_score,
        "has_website": lead.has_website,
        "has_ssl": lead.has_ssl,
        "has_email": lead.has_email,
        "has_phone": lead.has_phone,
        "description": (lead.description or "")[:500],
        "profile_url": lead.profile_url or "",
    }


CSV_FIELDS = [
    "name", "source_platform", "website", "email", "phone",
    "linkedin", "facebook", "instagram", "twitter", "github", "youtube",
    "industry", "city", "state", "country", "employee_estimate",
    "lead_score", "digital_score", "confidence_score",
    "has_website", "has_ssl", "has_email", "has_phone",
    "description", "profile_url",
]


@router.get("/{search_id}/exports/csv")
async def export_leads_csv(
    search_id: int,
    session: AsyncSession = Depends(get_session),
) -> Response:
    result = await session.execute(
        select(SocialLead).where(SocialLead.search_id == search_id)
        .order_by(SocialLead.lead_score.desc())
    )
    leads = result.scalars().all()

    output = io.StringIO()
    writer = csv.DictWriter(output, fieldnames=CSV_FIELDS)
    writer.writeheader()
    for lead in leads:
        writer.writerow(_build_lead_row(lead))

    content = output.getvalue()
    return Response(
        content=content,
        media_type="text/csv",
        headers={"Content-Disposition": f'attachment; filename="social-leads-{search_id}.csv"'},
    )


@router.get("/{search_id}/exports/json")
async def export_leads_json(
    search_id: int,
    session: AsyncSession = Depends(get_session),
) -> Response:
    result = await session.execute(
        select(SocialLead).where(SocialLead.search_id == search_id)
        .order_by(SocialLead.lead_score.desc())
    )
    leads = result.scalars().all()
    data = [_build_lead_row(lead) for lead in leads]
    content = json.dumps(data, indent=2, default=str)
    return Response(
        content=content,
        media_type="application/json",
        headers={"Content-Disposition": f'attachment; filename="social-leads-{search_id}.json"'},
    )


@router.get("/{search_id}/exports/excel")
async def export_leads_excel(
    search_id: int,
    session: AsyncSession = Depends(get_session),
) -> Response:
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment
    except ImportError:
        raise HTTPException(status_code=500, detail="openpyxl not installed")

    result = await session.execute(
        select(SocialLead).where(SocialLead.search_id == search_id)
        .order_by(SocialLead.lead_score.desc())
    )
    leads = result.scalars().all()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"Social Leads {search_id}"

    # Header styling
    header_fill = PatternFill(start_color="1A1A2E", end_color="1A1A2E", fill_type="solid")
    header_font = Font(bold=True, color="C0C1FF")

    for col_idx, field in enumerate(CSV_FIELDS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=field.replace("_", " ").title())
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    for row_idx, lead in enumerate(leads, start=2):
        row_data = _build_lead_row(lead)
        for col_idx, field in enumerate(CSV_FIELDS, start=1):
            ws.cell(row=row_idx, column=col_idx, value=row_data[field])

    # Auto-size columns
    for col in ws.columns:
        max_len = max((len(str(cell.value or "")) for cell in col), default=10)
        ws.column_dimensions[col[0].column_letter].width = min(max_len + 2, 40)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    return Response(
        content=buffer.read(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="social-leads-{search_id}.xlsx"'},
    )


@router.delete("/{search_id}", status_code=status.HTTP_204_NO_CONTENT)
async def delete_social_search(
    search_id: int,
    session: AsyncSession = Depends(get_session),
) -> None:
    search = await session.get(SocialSearch, search_id)
    if not search:
        raise HTTPException(status_code=404, detail="Search not found")
    await session.delete(search)
    await session.commit()


# ── Saved Searches ────────────────────────────────────────────────────────────

saved_router = APIRouter(prefix="/saved-searches", tags=["social"])


@saved_router.post("", response_model=SavedSearchRead, status_code=status.HTTP_201_CREATED)
async def create_saved_search(
    payload: SavedSearchCreate,
    session: AsyncSession = Depends(get_session),
) -> SavedSearchRead:
    saved = SavedSearch(
        name=payload.name,
        industry=payload.industry,
        location=payload.location,
        keywords=payload.keywords,
        sources=payload.sources,
        filters=payload.filters,
    )
    session.add(saved)
    await session.commit()
    await session.refresh(saved)
    return SavedSearchRead.model_validate(saved, from_attributes=True)


@saved_router.get("", response_model=list[SavedSearchRead])
async def list_saved_searches(
    session: AsyncSession = Depends(get_session),
) -> list[SavedSearchRead]:
    result = await session.execute(
        select(SavedSearch).order_by(SavedSearch.created_at.desc()).limit(50)
    )
    items = result.scalars().all()
    return [SavedSearchRead.model_validate(s, from_attributes=True) for s in items]


@saved_router.delete("/{saved_id}", status_code=status.HTTP_204_NO_CONTENT)
async def delete_saved_search(
    saved_id: int,
    session: AsyncSession = Depends(get_session),
) -> None:
    saved = await session.get(SavedSearch, saved_id)
    if not saved:
        raise HTTPException(status_code=404, detail="Saved search not found")
    await session.delete(saved)
    await session.commit()
