from __future__ import annotations

import csv
from io import BytesIO, StringIO

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy.orm import selectinload

from app.models.entities import Business


class ExportService:
    async def build_rows(self, session: AsyncSession, job_id: int) -> list[dict]:
        statement = (
            select(Business)
            .where(Business.job_id == job_id)
            .order_by(Business.id.asc())
            .options(
                selectinload(Business.emails),
                selectinload(Business.social_accounts),
                selectinload(Business.website_record),
                selectinload(Business.lead_score),
            )
        )
        result = await session.execute(statement)
        rows = []
        for business in result.scalars().unique().all():
            # Emails
            primary_email = next((e.address for e in business.emails if e.is_primary), "")
            all_emails = "; ".join(e.address for e in business.emails)

            # Phones: Maps-scraped + website-extracted extras
            maps_phone = business.phone_number or ""
            website_phones = (business.raw_payload or {}).get("website_phones", [])
            all_phones = "; ".join(filter(None, [maps_phone] + [p for p in website_phones if p != maps_phone]))

            # Social links grouped by platform
            socials: dict[str, list[str]] = {}
            for sa in business.social_accounts:
                socials.setdefault(sa.platform, []).append(sa.url)

            # Website analysis
            ws = business.website_record

            # Lead scores
            ls = business.lead_score

            rows.append(
                {
                    # Core identity
                    "name": business.name,
                    "category": business.category or "",
                    "address": business.address or "",
                    "business_status": business.business_status or "",
                    # Contact details
                    "phone_number": maps_phone,
                    "all_phones": all_phones,
                    "primary_email": primary_email,
                    "all_emails": all_emails,
                    # Website
                    "website": business.website or "",
                    "final_url": ws.final_url if ws else "",
                    "ssl_enabled": ws.ssl_enabled if ws else "",
                    "mobile_friendly": ws.mobile_viewport if ws else "",
                    # SEO
                    "meta_title": ws.meta_title if ws else "",
                    "meta_description": ws.meta_description if ws else "",
                    "has_contact_page": ws.contact_page_found if ws else "",
                    "has_google_analytics": ws.google_analytics if ws else "",
                    "has_facebook_pixel": ws.facebook_pixel if ws else "",
                    # Social profiles
                    "facebook": "; ".join(socials.get("facebook", [])),
                    "instagram": "; ".join(socials.get("instagram", [])),
                    "linkedin": "; ".join(socials.get("linkedin", [])),
                    "twitter_x": "; ".join(socials.get("twitter", [])),
                    "youtube": "; ".join(socials.get("youtube", [])),
                    "whatsapp": "; ".join(socials.get("whatsapp", [])),
                    # Ratings
                    "rating": business.rating or "",
                    "review_count": business.review_count or "",
                    "opening_hours": business.opening_hours or "",
                    # Lead scores
                    "digital_presence_score": ls.digital_presence_score if ls else "",
                    "lead_opportunity_score": ls.lead_opportunity_score if ls else "",
                    # Geo & Maps
                    "latitude": business.latitude or "",
                    "longitude": business.longitude or "",
                    "google_maps_url": business.google_maps_url,
                }
            )
        return rows

    def to_csv(self, rows: list[dict]) -> bytes:
        output = StringIO()
        writer = csv.DictWriter(output, fieldnames=list(rows[0].keys()) if rows else [])
        if rows:
            writer.writeheader()
            writer.writerows(rows)
        return output.getvalue().encode("utf-8")

    def to_excel(self, rows: list[dict]) -> bytes:
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "Leads"
        if rows:
            headers = list(rows[0].keys())
            # Write header row with styling
            for col_idx, header in enumerate(headers, start=1):
                cell = worksheet.cell(row=1, column=col_idx, value=header.replace("_", " ").title())
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill(start_color="1E3A5F", end_color="1E3A5F", fill_type="solid")
            # Write data rows
            for row_idx, row in enumerate(rows, start=2):
                for col_idx, value in enumerate(row.values(), start=1):
                    worksheet.cell(row=row_idx, column=col_idx, value=str(value) if isinstance(value, bool) else value)
            # Auto-size columns (approximate)
            for col in worksheet.columns:
                max_len = max((len(str(cell.value or "")) for cell in col), default=10)
                worksheet.column_dimensions[col[0].column_letter].width = min(max_len + 2, 60)
        output = BytesIO()
        workbook.save(output)
        return output.getvalue()

